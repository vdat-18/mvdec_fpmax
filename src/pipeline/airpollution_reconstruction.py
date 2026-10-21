"""Clean-room reconstruction of the paper's Luxembourg feature matrix.

This module rebuilds the spatial grid and traffic-distance features from official
raw data. It intentionally keeps competing undocumented protocol choices visible
instead of selecting a variant only because it resembles the published score.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.spatial import cKDTree

PAPER_ROW_COUNT = 3765
PAPER_GRID_CRS = "EPSG:4326"
DISTANCE_CRS = "EPSG:2169"
DEFAULT_RESOLUTIONS = (0.01, 0.00926)
TRAFFIC_LOW_THRESHOLD = 120_000.0
TRAFFIC_HIGH_THRESHOLD = 600_000.0
TRAFFIC_MODES = ("dir3_u", "dir12_u", "all_u", "all_rows")
TRAFFIC_FEATURES = {
    "high": "high_traffic",
    "medium": "medium_traffic",
    "low": "low_traffic",
}
ROAD_CANDIDATES = {
    "osm_primary": ("primary",),
    "osm_motorway": ("motorway",),
    "osm_motorway_trunk": ("motorway", "trunk"),
}


@dataclass(frozen=True)
class GridMetadata:
    """Reconstructed regular-grid geometry and row-count diagnostics."""

    resolution_degrees: float
    longitude_bins: int
    latitude_bins: int
    bounding_box_cells: int
    retained_cells: int
    paper_row_gap: int
    bounds: list[float]


@dataclass(frozen=True)
class CandidateDiagnostic:
    """Comparison of one grid and traffic aggregation protocol."""

    resolution_degrees: float
    traffic_mode: str
    landuse_source: str
    rows: int
    station_category_counts: dict[str, int]
    traffic_quantile_rmse: dict[str, float]
    traffic_quantile_rmse_mean: float | None
    road_quantile_rmse: dict[str, float]
    landuse_quantile_rmse: dict[str, float]
    output_csv: str


def default_raw_root() -> Path:
    return (
        Path(__file__).resolve().parents[2]
        / "data"
        / "public"
        / "raw"
        / "airpollution_reconstruction"
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _require_file(path: Path, label: str) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"Missing {label}: {path}")


def classify_traffic(value: float) -> str:
    """Apply the two thresholds stated on page 11 of the paper."""

    if not np.isfinite(value) or value < 0:
        raise ValueError("Traffic values must be finite and non-negative.")
    if value >= TRAFFIC_HIGH_THRESHOLD:
        return "high"
    if value >= TRAFFIC_LOW_THRESHOLD:
        return "medium"
    return "low"


def minmax(values: np.ndarray) -> np.ndarray:
    """Scale one finite vector to [0, 1]."""

    array = np.asarray(values, dtype=np.float64)
    if array.ndim != 1 or len(array) == 0 or not np.isfinite(array).all():
        raise ValueError("Min-Max input must be a non-empty finite vector.")
    span = float(array.max() - array.min())
    if span == 0.0:
        return np.zeros_like(array)
    return (array - array.min()) / span


def extract_traffic_station_variants(workbook_path: Path) -> pd.DataFrame:
    """Extract all defensible interpretations of the official 2018 workbook.

    The workbook stores vehicle classes and direction totals in overlapping rows.
    The paper does not say which rows were aggregated, so all four choices remain
    explicit in the output.
    """

    from openpyxl import load_workbook
    from pyproj import Transformer

    _require_file(workbook_path, "traffic workbook")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Data" not in workbook.sheetnames:
        raise ValueError("Traffic workbook does not contain the expected Data sheet.")
    sheet = workbook["Data"]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    required = {
        "POSTE_ID",
        "DATECOM",
        "DIRECTION",
        "VEHICULE",
        "SUM_TRAF",
        "LOCALITE",
        "ROUTE",
        "COORD_X",
        "COORD_Y",
    }
    missing = required - set(header)
    if missing:
        raise ValueError(f"Traffic workbook is missing columns: {sorted(missing)}")
    positions = {name: header.index(name) for name in required}

    monthly: dict[str, defaultdict[tuple[int, int, int], float]] = {
        mode: defaultdict(float) for mode in TRAFFIC_MODES
    }
    metadata: dict[int, tuple[float, float, str, str]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        station = row[positions["POSTE_ID"]]
        date = row[positions["DATECOM"]]
        direction = row[positions["DIRECTION"]]
        vehicle = row[positions["VEHICULE"]]
        total = row[positions["SUM_TRAF"]]
        x_coord = row[positions["COORD_X"]]
        y_coord = row[positions["COORD_Y"]]
        if None in (station, date, total, x_coord, y_coord):
            continue
        station_id = int(station)
        key = (station_id, int(date.year), int(date.month))
        total_value = float(total)
        metadata[station_id] = (
            float(x_coord),
            float(y_coord),
            str(row[positions["LOCALITE"]]),
            str(row[positions["ROUTE"]]),
        )
        monthly["all_rows"][key] += total_value
        if vehicle == "U":
            monthly["all_u"][key] += total_value
            if direction == 3:
                monthly["dir3_u"][key] += total_value
            if direction in (1, 2):
                monthly["dir12_u"][key] += total_value

    transformer = Transformer.from_crs(DISTANCE_CRS, PAPER_GRID_CRS, always_xy=True)
    records: list[dict[str, object]] = []
    for station_id, (x_coord, y_coord, locality, route) in sorted(metadata.items()):
        longitude, latitude = transformer.transform(x_coord, y_coord)
        record: dict[str, object] = {
            "station_id": station_id,
            "x_luref": x_coord,
            "y_luref": y_coord,
            "longitude": longitude,
            "latitude": latitude,
            "locality": locality,
            "route": route,
        }
        for mode in TRAFFIC_MODES:
            values = [
                value
                for (source_station, _, _), value in monthly[mode].items()
                if source_station == station_id
            ]
            if not values:
                raise ValueError(
                    f"Traffic mode {mode!r} has no values for station {station_id}."
                )
            mean_monthly = float(np.mean(values))
            record[f"{mode}_mean_monthly"] = mean_monthly
            record[f"{mode}_observed_months"] = len(values)
            record[f"{mode}_category"] = classify_traffic(mean_monthly)
        records.append(record)
    if not records:
        raise ValueError("No valid traffic stations were extracted.")
    return pd.DataFrame.from_records(records)


def load_country_boundary(boundary_path: Path):
    """Load and union the official Luxembourg commune boundaries."""

    import geopandas as gpd

    _require_file(boundary_path, "Luxembourg boundary")
    frame = gpd.read_file(boundary_path)
    if frame.empty or frame.crs is None:
        raise ValueError("Boundary file must contain georeferenced geometries.")
    boundary = frame.to_crs(PAPER_GRID_CRS).geometry.union_all()
    if boundary.is_empty or not boundary.is_valid:
        raise ValueError("Luxembourg boundary is empty or invalid.")
    return boundary


def build_regular_grid(
    boundary, resolution: float
) -> tuple[pd.DataFrame, GridMetadata]:
    """Recreate the S5P-tools bbox-aligned grid and retain centers in Luxembourg."""

    import geopandas as gpd
    from shapely import contains_xy

    if not np.isfinite(resolution) or resolution <= 0:
        raise ValueError("Grid resolution must be a finite positive number.")
    min_x, min_y, max_x, max_y = boundary.bounds
    longitude_edges = int(abs(max_x - min_x) / resolution + 1)
    latitude_edges = int(abs(max_y - min_y) / resolution + 1)
    longitude_bins = longitude_edges - 1
    latitude_bins = latitude_edges - 1
    longitudes = min_x + (np.arange(longitude_bins) + 0.5) * resolution
    latitudes = min_y + (np.arange(latitude_bins) + 0.5) * resolution
    longitude_grid, latitude_grid = np.meshgrid(longitudes, latitudes)
    row_grid, column_grid = np.indices((latitude_bins, longitude_bins))
    mask = contains_xy(
        boundary,
        longitude_grid.ravel(),
        latitude_grid.ravel(),
    )
    retained_longitudes = longitude_grid.ravel()[mask]
    retained_latitudes = latitude_grid.ravel()[mask]
    points = gpd.GeoSeries(
        gpd.points_from_xy(retained_longitudes, retained_latitudes),
        crs=PAPER_GRID_CRS,
    ).to_crs(DISTANCE_CRS)
    grid = pd.DataFrame(
        {
            "grid_index": np.arange(len(points), dtype=np.int64),
            "grid_row": row_grid.ravel()[mask],
            "grid_column": column_grid.ravel()[mask],
            "longitude": retained_longitudes,
            "latitude": retained_latitudes,
            "x_luref": points.x.to_numpy(),
            "y_luref": points.y.to_numpy(),
        }
    )
    metadata = GridMetadata(
        resolution_degrees=float(resolution),
        longitude_bins=longitude_bins,
        latitude_bins=latitude_bins,
        bounding_box_cells=longitude_bins * latitude_bins,
        retained_cells=len(grid),
        paper_row_gap=len(grid) - PAPER_ROW_COUNT,
        bounds=[float(value) for value in boundary.bounds],
    )
    return grid, metadata


def add_traffic_distance_features(
    grid: pd.DataFrame,
    stations: pd.DataFrame,
    mode: str,
) -> tuple[pd.DataFrame, dict[str, int]]:
    """Compute nearest-sensor distances and the paper-like [0, 1] columns."""

    if mode not in TRAFFIC_MODES:
        raise ValueError(f"Unsupported traffic aggregation mode: {mode!r}")
    output = grid.copy()
    grid_coordinates = output[["x_luref", "y_luref"]].to_numpy()
    category_column = f"{mode}_category"
    category_counts: dict[str, int] = {}
    for category, feature_name in TRAFFIC_FEATURES.items():
        selected = stations.loc[
            stations[category_column] == category,
            ["x_luref", "y_luref"],
        ].to_numpy()
        if len(selected) == 0:
            raise ValueError(f"Traffic mode {mode!r} has no {category} stations.")
        distances = cKDTree(selected).query(grid_coordinates)[0]
        output[f"{feature_name}_meters"] = distances
        output[feature_name] = minmax(distances)
        category_counts[category] = int(len(selected))
    return output, category_counts


def load_road_geometries(osm_json_path: Path, boundary) -> dict[str, object]:
    """Load road-line candidates from a historical Overpass JSON export."""

    import geopandas as gpd
    from shapely.geometry import LineString

    _require_file(osm_json_path, "historical OSM road export")
    payload = json.loads(osm_json_path.read_text(encoding="utf-8"))
    if "elements" not in payload:
        raise ValueError("Overpass JSON does not contain an elements array.")
    tagged_lines: dict[str, list[object]] = defaultdict(list)
    for element in payload["elements"]:
        highway = element.get("tags", {}).get("highway")
        geometry = element.get("geometry")
        if highway is None or not geometry:
            continue
        coordinates = [(point["lon"], point["lat"]) for point in geometry]
        if len(coordinates) < 2:
            continue
        line = LineString(coordinates)
        if line.intersects(boundary):
            tagged_lines[str(highway)].append(line.intersection(boundary))

    candidates: dict[str, object] = {}
    for name, highway_tags in ROAD_CANDIDATES.items():
        lines = [
            line
            for highway_tag in highway_tags
            for line in tagged_lines.get(highway_tag, [])
        ]
        if not lines:
            continue
        candidates[name] = (
            gpd.GeoSeries(lines, crs=PAPER_GRID_CRS).to_crs(DISTANCE_CRS).union_all()
        )
    if not candidates:
        raise ValueError("No supported highway geometries intersect Luxembourg.")
    return candidates


def add_road_distance_candidates(
    grid: pd.DataFrame,
    road_geometries: dict[str, object],
) -> pd.DataFrame:
    """Add raw meter distances for each defensible OSM road interpretation."""

    import geopandas as gpd
    from shapely import distance

    output = grid.copy()
    points = gpd.GeoSeries(
        gpd.points_from_xy(output["x_luref"], output["y_luref"]),
        crs=DISTANCE_CRS,
    )
    for name, geometry in road_geometries.items():
        output[f"highway_primary_{name}"] = distance(points.array, geometry)
    return output


def load_landuse_geometries(osm_json_path: Path, boundary) -> dict[str, object]:
    """Load polygonal OSM land-use candidates from a historical ways export."""

    import geopandas as gpd
    from shapely import make_valid
    from shapely.geometry import Polygon

    _require_file(osm_json_path, "historical OSM land-use export")
    payload = json.loads(osm_json_path.read_text(encoding="utf-8"))
    classes: dict[str, list[object]] = defaultdict(list)
    for element in payload.get("elements", []):
        tags = element.get("tags", {})
        landuse = tags.get("landuse")
        natural = tags.get("natural")
        category = None
        if landuse in {"industrial", "residential", "forest"}:
            category = landuse
        elif natural == "wood":
            category = "forest"
        geometry = element.get("geometry")
        if category is None or not geometry:
            continue
        coordinates = [(point["lon"], point["lat"]) for point in geometry]
        if len(coordinates) < 4 or coordinates[0] != coordinates[-1]:
            continue
        polygon = make_valid(Polygon(coordinates))
        if polygon.is_empty or not polygon.intersects(boundary):
            continue
        classes[category].append(polygon.intersection(boundary))

    output: dict[str, object] = {}
    for category in ("industrial", "residential", "forest"):
        polygons = classes.get(category, [])
        if polygons:
            output[category] = (
                gpd.GeoSeries(
                    polygons,
                    crs=PAPER_GRID_CRS,
                )
                .to_crs(DISTANCE_CRS)
                .union_all()
            )
    if set(output) != {"industrial", "residential", "forest"}:
        raise ValueError(
            "Historical OSM export does not contain all three land-use classes."
        )
    return output


def load_geofabrik_landuse_geometries(
    shapefile_zip_path: Path,
    boundary,
) -> dict[str, object]:
    """Load relation-complete land-use polygons from a Geofabrik snapshot."""

    import geopandas as gpd

    _require_file(shapefile_zip_path, "historical Geofabrik shapefile archive")
    archive_uri = f"zip://{shapefile_zip_path.resolve()}!gis_osm_landuse_a_free_1.shp"
    frame = gpd.read_file(archive_uri)
    if "fclass" not in frame.columns or frame.crs is None:
        raise ValueError(
            "Geofabrik land-use layer must provide fclass values and a CRS."
        )
    frame = frame.to_crs(PAPER_GRID_CRS)
    frame = frame[frame["fclass"].isin({"industrial", "residential", "forest"})]
    frame = frame.loc[frame.geometry.notna() & ~frame.geometry.is_empty].copy()
    frame["geometry"] = frame.geometry.make_valid().intersection(boundary)
    frame = frame.loc[~frame.geometry.is_empty]

    output: dict[str, object] = {}
    for category in ("industrial", "residential", "forest"):
        geometries = frame.loc[frame["fclass"] == category, "geometry"]
        if not geometries.empty:
            output[category] = (
                gpd.GeoSeries(geometries, crs=PAPER_GRID_CRS)
                .to_crs(DISTANCE_CRS)
                .union_all()
            )
    if set(output) != {"industrial", "residential", "forest"}:
        raise ValueError(
            "Geofabrik snapshot does not contain all three land-use classes."
        )
    return output


def _queen_neighbor_mean(grid: pd.DataFrame, values: np.ndarray) -> np.ndarray:
    lookup = {
        (int(row), int(column)): float(value)
        for row, column, value in zip(
            grid["grid_row"],
            grid["grid_column"],
            values,
            strict=True,
        )
    }
    means = np.empty(len(grid), dtype=np.float64)
    for index, (row, column) in enumerate(
        zip(grid["grid_row"], grid["grid_column"], strict=True)
    ):
        neighbors = [
            lookup[(int(row) + row_offset, int(column) + column_offset)]
            for row_offset in (-1, 0, 1)
            for column_offset in (-1, 0, 1)
            if (row_offset != 0 or column_offset != 0)
            and (int(row) + row_offset, int(column) + column_offset) in lookup
        ]
        means[index] = float(np.mean(neighbors)) if neighbors else 0.0
    return means


def _grid_disk_neighbor_mean(
    grid: pd.DataFrame,
    values: np.ndarray,
    radius_cells: float,
) -> tuple[np.ndarray, np.ndarray]:
    """Apply a fixed index-space disk, clipped only by the retained grid mask."""

    if radius_cells <= 0:
        raise ValueError("Grid-disk radius must be positive.")
    lookup = {
        (int(row), int(column)): float(value)
        for row, column, value in zip(
            grid["grid_row"],
            grid["grid_column"],
            values,
            strict=True,
        )
    }
    extent = int(np.floor(radius_cells))
    offsets = [
        (row_offset, column_offset)
        for row_offset in range(-extent, extent + 1)
        for column_offset in range(-extent, extent + 1)
        if row_offset**2 + column_offset**2 <= radius_cells**2
    ]
    means = np.empty(len(grid), dtype=np.float64)
    supports = np.empty(len(grid), dtype=np.int64)
    for index, (row, column) in enumerate(
        zip(grid["grid_row"], grid["grid_column"], strict=True)
    ):
        neighborhood = [
            lookup[(int(row) + row_offset, int(column) + column_offset)]
            for row_offset, column_offset in offsets
            if (int(row) + row_offset, int(column) + column_offset) in lookup
        ]
        supports[index] = len(neighborhood)
        means[index] = float(np.mean(neighborhood))
    return means, supports


def _radius_neighbor_mean(
    coordinates: np.ndarray,
    values: np.ndarray,
    radius_meters: float,
) -> np.ndarray:
    tree = cKDTree(coordinates)
    neighborhoods = tree.query_ball_point(coordinates, radius_meters)
    return np.fromiter(
        (float(np.mean(values[indices])) for indices in neighborhoods),
        dtype=np.float64,
        count=len(coordinates),
    )


def add_landuse_candidates(
    grid: pd.DataFrame,
    landuse_geometries: dict[str, object],
) -> pd.DataFrame:
    """Add explicit land-use raster and neighborhood hypotheses."""

    import geopandas as gpd
    from shapely import intersects

    output = grid.copy()
    points = gpd.GeoSeries(
        gpd.points_from_xy(output["x_luref"], output["y_luref"]),
        crs=DISTANCE_CRS,
    )
    coordinates = output[["x_luref", "y_luref"]].to_numpy()
    for category, geometry in landuse_geometries.items():
        center = intersects(points.array, geometry).astype(np.float64)
        queen8 = _queen_neighbor_mean(output, center)
        disk2p5, disk2p5_support = _grid_disk_neighbor_mean(output, center, 2.5)
        radius_2500 = _radius_neighbor_mean(coordinates, center, 2_500.0)
        radius_5000 = _radius_neighbor_mean(coordinates, center, 5_000.0)
        prefix = f"landuse_{category}"
        output[f"{prefix}_center"] = center
        output[f"{prefix}_queen8_mean"] = queen8
        output[f"{prefix}_disk2p5_mean"] = disk2p5
        output[f"{prefix}_disk2p5_support"] = disk2p5_support
        output[f"{prefix}_radius2500_mean"] = radius_2500
        output[f"{prefix}_radius5000_mean"] = radius_5000
        if category == "industrial":
            output[f"{prefix}_disk2p5_variance"] = disk2p5 * (1.0 - disk2p5)
            output[f"{prefix}_radius2500_variance"] = radius_2500 * (1.0 - radius_2500)
            output[f"{prefix}_radius5000_variance"] = radius_5000 * (1.0 - radius_5000)
    return output


def quantile_rmse(reference: np.ndarray, candidate: np.ndarray) -> float:
    """Compare distributions without assuming row correspondence."""

    reference_array = np.asarray(reference, dtype=np.float64)
    candidate_array = np.asarray(candidate, dtype=np.float64)
    if (
        reference_array.ndim != 1
        or candidate_array.ndim != 1
        or len(reference_array) == 0
        or len(candidate_array) == 0
        or not np.isfinite(reference_array).all()
        or not np.isfinite(candidate_array).all()
    ):
        raise ValueError("Quantile RMSE requires non-empty finite vectors.")
    probabilities = np.linspace(0.0, 1.0, 1001)
    difference = np.quantile(reference_array, probabilities) - np.quantile(
        candidate_array, probabilities
    )
    return float(np.sqrt(np.mean(np.square(difference))))


def _compare_candidate(
    reconstructed: pd.DataFrame,
    author_frame: pd.DataFrame | None,
) -> tuple[dict[str, float], dict[str, float], dict[str, float]]:
    if author_frame is None:
        return {}, {}, {}

    def require_author_columns(columns: tuple[str, ...], family: str) -> None:
        missing = [column for column in columns if column not in author_frame.columns]
        if missing:
            missing_list = ", ".join(missing)
            raise ValueError(
                f"Author data is missing required {family} columns: {missing_list}."
            )

    traffic_columns = tuple(TRAFFIC_FEATURES.values())
    present_traffic = [
        column for column in traffic_columns if column in reconstructed.columns
    ]
    if present_traffic and len(present_traffic) != len(traffic_columns):
        missing = [
            column for column in traffic_columns if column not in reconstructed.columns
        ]
        raise ValueError(
            "Traffic candidate must contain the complete feature family; missing: "
            + ", ".join(missing)
            + "."
        )

    traffic_rmse: dict[str, float] = {}
    if present_traffic:
        require_author_columns(traffic_columns, "traffic")
        traffic_rmse = {
            feature_name: quantile_rmse(
                author_frame[feature_name].to_numpy(),
                reconstructed[feature_name].to_numpy(),
            )
            for feature_name in traffic_columns
        }

    road_columns = [
        column
        for column in reconstructed.columns
        if column.startswith("highway_primary_")
    ]
    road_rmse: dict[str, float] = {}
    if road_columns:
        require_author_columns(("highway_primary",), "road")
        road_rmse = {
            column.removeprefix("highway_primary_"): quantile_rmse(
                author_frame["highway_primary"].to_numpy(),
                reconstructed[column].to_numpy(),
            )
            for column in road_columns
        }

    landuse_targets = {
        "industrial": "landuse_industrial_b5",
        "residential": "landuse_residential_b5",
        "forest": "landuse_forest_b5",
    }
    landuse_rmse: dict[str, float] = {}
    for category, target in landuse_targets.items():
        for hypothesis in (
            "queen8_mean",
            "disk2p5_mean",
            "radius2500_mean",
            "radius5000_mean",
        ):
            column = f"landuse_{category}_{hypothesis}"
            if column not in reconstructed.columns:
                continue
            require_author_columns((target,), f"land-use {category}")
            landuse_rmse[f"{category}_{hypothesis}_to_b5"] = quantile_rmse(
                author_frame[target].to_numpy(),
                reconstructed[column].to_numpy(),
            )

    industrial_targets = {
        "mean": "landuse_industrial_distance_mean",
        "variance": "landuse_industrial_variance",
    }
    for hypothesis in ("disk2p5", "radius2500", "radius5000"):
        for statistic, target in industrial_targets.items():
            column = f"landuse_industrial_{hypothesis}_{statistic}"
            if column not in reconstructed.columns:
                continue
            require_author_columns((target,), f"land-use industrial {statistic}")
            landuse_rmse[f"industrial_{hypothesis}_{statistic}"] = quantile_rmse(
                author_frame[target].to_numpy(),
                reconstructed[column].to_numpy(),
            )

    return traffic_rmse, road_rmse, landuse_rmse


def _summarize_landuse_matches(
    diagnostics: list[dict[str, object]],
) -> list[dict[str, object]]:
    """Select the best kernel per target without duplicating traffic variants."""

    target_filters = {
        "industrial_b5": lambda key: (
            key.startswith("industrial_") and key.endswith("_to_b5")
        ),
        "residential_b5": lambda key: (
            key.startswith("residential_") and key.endswith("_to_b5")
        ),
        "forest_b5": lambda key: key.startswith("forest_") and key.endswith("_to_b5"),
        "industrial_distance_mean": lambda key: (
            key.startswith("industrial_")
            and key.endswith("_mean")
            and not key.endswith("_to_b5")
        ),
        "industrial_variance": lambda key: (
            key.startswith("industrial_") and key.endswith("_variance")
        ),
    }
    summaries: list[dict[str, object]] = []
    seen: set[tuple[float, str]] = set()
    for diagnostic in diagnostics:
        resolution = float(diagnostic["resolution_degrees"])
        landuse_source = str(diagnostic["landuse_source"])
        source_key = (resolution, landuse_source)
        if source_key in seen:
            continue
        seen.add(source_key)
        values = diagnostic["landuse_quantile_rmse"]
        if not isinstance(values, dict):
            raise TypeError("Land-use diagnostics must be a dictionary.")
        for target, matches in target_filters.items():
            candidates = [
                (str(candidate), float(rmse))
                for candidate, rmse in values.items()
                if matches(str(candidate))
            ]
            if not candidates:
                continue
            candidate, rmse = min(candidates, key=lambda item: item[1])
            summaries.append(
                {
                    "resolution_degrees": resolution,
                    "landuse_source": landuse_source,
                    "target": target,
                    "candidate": candidate,
                    "quantile_rmse": rmse,
                }
            )
    return summaries


def reconstruct_stage1(
    *,
    boundary_path: Path,
    traffic_workbook_path: Path,
    output_dir: Path,
    resolutions: tuple[float, ...] = DEFAULT_RESOLUTIONS,
    author_csv_path: Path | None = None,
    osm_json_path: Path | None = None,
    landuse_osm_json_path: Path | None = None,
    landuse_geofabrik_zip_path: Path | None = None,
) -> dict[str, object]:
    """Reconstruct the grid, traffic features, and road-distance candidates."""

    if not resolutions:
        raise ValueError("At least one grid resolution is required.")
    boundary = load_country_boundary(boundary_path)
    stations = extract_traffic_station_variants(traffic_workbook_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    stations_path = output_dir / "traffic_stations_2018.csv"
    stations.to_csv(stations_path, index=False)

    author_frame: pd.DataFrame | None = None
    if author_csv_path is not None:
        from pipeline.airpollution_forensic import load_airpollution_csv

        author_frame = load_airpollution_csv(author_csv_path)

    road_geometries: dict[str, object] = {}
    if osm_json_path is not None:
        road_geometries = load_road_geometries(osm_json_path, boundary)
    landuse_sources: dict[str, dict[str, object]] = {}
    if landuse_osm_json_path is not None:
        landuse_sources["overpass_ways"] = load_landuse_geometries(
            landuse_osm_json_path,
            boundary,
        )
    if landuse_geofabrik_zip_path is not None:
        landuse_sources["geofabrik_2021_01_01"] = load_geofabrik_landuse_geometries(
            landuse_geofabrik_zip_path,
            boundary,
        )
    if not landuse_sources:
        landuse_sources["none"] = {}

    grid_metadata: list[dict[str, object]] = []
    diagnostics: list[CandidateDiagnostic] = []
    for resolution in resolutions:
        base_grid, metadata = build_regular_grid(boundary, resolution)
        grid_metadata.append(asdict(metadata))
        if road_geometries:
            base_grid = add_road_distance_candidates(base_grid, road_geometries)
        resolution_token = f"{resolution:.6f}".rstrip("0").rstrip(".")
        for landuse_source, landuse_geometries in landuse_sources.items():
            grid = base_grid
            if landuse_geometries:
                grid = add_landuse_candidates(grid, landuse_geometries)
            for mode in TRAFFIC_MODES:
                reconstructed, category_counts = add_traffic_distance_features(
                    grid,
                    stations,
                    mode,
                )
                output_csv = output_dir / (
                    f"grid_{resolution_token}_landuse_{landuse_source}_"
                    f"traffic_{mode}.csv"
                )
                reconstructed.to_csv(output_csv, index=False)
                traffic_rmse, road_rmse, landuse_rmse = _compare_candidate(
                    reconstructed,
                    author_frame,
                )
                diagnostics.append(
                    CandidateDiagnostic(
                        resolution_degrees=float(resolution),
                        traffic_mode=mode,
                        landuse_source=landuse_source,
                        rows=len(reconstructed),
                        station_category_counts=category_counts,
                        traffic_quantile_rmse=traffic_rmse,
                        traffic_quantile_rmse_mean=(
                            float(np.mean(list(traffic_rmse.values())))
                            if traffic_rmse
                            else None
                        ),
                        road_quantile_rmse=road_rmse,
                        landuse_quantile_rmse=landuse_rmse,
                        output_csv=str(output_csv.resolve()),
                    )
                )

    serialized_diagnostics = [asdict(item) for item in diagnostics]
    comparable = [
        item
        for item in serialized_diagnostics
        if item["traffic_quantile_rmse_mean"] is not None
    ]
    best_traffic = min(
        comparable,
        key=lambda item: item["traffic_quantile_rmse_mean"],
        default=None,
    )
    report = {
        "schema_version": 1,
        "status": "stage1_not_final_13_feature_matrix",
        "paper_target": {
            "rows": PAPER_ROW_COUNT,
            "features": 13,
            "stated_grid_crs": PAPER_GRID_CRS,
            "s5p_tools_default_resolution_degrees": 0.01,
            "traffic_thresholds": {
                "low_to_medium": TRAFFIC_LOW_THRESHOLD,
                "medium_to_high": TRAFFIC_HIGH_THRESHOLD,
            },
        },
        "inputs": {
            "boundary": {
                "path": str(boundary_path.resolve()),
                "sha256": _sha256(boundary_path),
            },
            "traffic_workbook": {
                "path": str(traffic_workbook_path.resolve()),
                "sha256": _sha256(traffic_workbook_path),
            },
            "author_csv": (
                {
                    "path": str(author_csv_path.resolve()),
                    "sha256": _sha256(author_csv_path),
                }
                if author_csv_path is not None
                else None
            ),
            "historical_osm": (
                {
                    "path": str(osm_json_path.resolve()),
                    "sha256": _sha256(osm_json_path),
                }
                if osm_json_path is not None
                else None
            ),
            "historical_osm_landuse": (
                {
                    "path": str(landuse_osm_json_path.resolve()),
                    "sha256": _sha256(landuse_osm_json_path),
                }
                if landuse_osm_json_path is not None
                else None
            ),
            "historical_geofabrik_landuse": (
                {
                    "path": str(landuse_geofabrik_zip_path.resolve()),
                    "sha256": _sha256(landuse_geofabrik_zip_path),
                }
                if landuse_geofabrik_zip_path is not None
                else None
            ),
        },
        "traffic_station_rows": len(stations),
        "traffic_modes": list(TRAFFIC_MODES),
        "grid_candidates": grid_metadata,
        "candidate_diagnostics": serialized_diagnostics,
        "best_author_traffic_distribution_match": best_traffic,
        "best_landuse_distribution_matches": _summarize_landuse_matches(
            serialized_diagnostics
        ),
        "interpretation": [
            (
                "The paper does not specify which overlapping direction/vehicle "
                "rows in the official workbook were aggregated."
            ),
            (
                "The S5P-tools default 0.01-degree grid produces a different row "
                "count from the paper, so the exact grid/mask is still unresolved."
            ),
            (
                "The fixed 2.5-cell disk has 21 interior cells and is retained as "
                "an explicit hypothesis because the author CSV contains many "
                "fractions with denominator 21; this is not proof of the kernel."
            ),
            (
                "A low distribution gap is forensic evidence only; it does not "
                "establish exact row correspondence or justify post-hoc selection."
            ),
        ],
        "remaining_stages": [
            "Resolve the exact grid origin, mask, and 3,765 retained cells.",
            "Resolve the exact land-use source and convolution kernel.",
            "Acquire and aggregate Sentinel-5P NO2 from 2018 through June 2021.",
            "Resolve the undocumented WRF-Chem column-to-ground conversion.",
            "Rebuild neighbor-NO2 features and validate against Table 4.",
        ],
    }
    (output_dir / "stage1_report.json").write_text(
        json.dumps(report, indent=2, sort_keys=True, allow_nan=False),
        encoding="utf-8",
    )
    write_stage1_markdown(report, output_dir / "stage1_report.md")
    return report


def write_stage1_markdown(report: dict[str, object], path: Path) -> None:
    """Write the human-readable reconstruction status report."""

    best = report["best_author_traffic_distribution_match"]
    lines = [
        "# Air Pollution Clean-room Reconstruction - Stage 1",
        "",
        "## Status",
        "",
        "This is a raw-source reconstruction, not a transformed copy of the "
        "author CSV. Stage 1 rebuilds the spatial grid, official traffic station "
        "aggregation variants, nearest-category distances, historical OSM "
        "road-distance candidates, and explicit Overpass/Geofabrik land-use "
        "alternatives.",
        "",
        f"- Official traffic stations recovered: `{report['traffic_station_rows']}`",
        f"- Paper target rows: `{report['paper_target']['rows']}`",
        "- Final 13-feature matrix complete: `False`",
        "",
        "## Grid candidates",
        "",
        "| Resolution | Retained cells | Gap to 3,765 | Bbox cells |",
        "|---:|---:|---:|---:|",
    ]
    for grid in report["grid_candidates"]:
        lines.append(
            f"| {grid['resolution_degrees']:.6f} | {grid['retained_cells']} | "
            f"{grid['paper_row_gap']:+d} | {grid['bounding_box_cells']} |"
        )
    lines.extend(
        [
            "",
            "## Traffic protocol diagnostics",
            "",
            "| Resolution | Land-use source | Workbook aggregation | "
            "High/medium/low stations | Mean quantile RMSE |",
            "|---:|---|---|---|---:|",
        ]
    )
    for item in report["candidate_diagnostics"]:
        counts = item["station_category_counts"]
        value = item["traffic_quantile_rmse_mean"]
        display_value = f"{value:.6f}" if value is not None else "n/a"
        lines.append(
            f"| {item['resolution_degrees']:.6f} | "
            f"`{item['landuse_source']}` | `{item['traffic_mode']}` | "
            f"{counts['high']}/{counts['medium']}/{counts['low']} | "
            f"{display_value} |"
        )
    if best is not None:
        lines.extend(
            [
                "",
                "The closest traffic-distribution diagnostic is "
                f"`{best['traffic_mode']}` at resolution "
                f"`{best['resolution_degrees']:.6f}` with mean RMSE "
                f"`{best['traffic_quantile_rmse_mean']:.6f}`. This is not yet "
                "accepted as the paper protocol because row correspondence and "
                "the remaining features are unresolved.",
            ]
        )
    lines.extend(
        [
            "",
            "## Land-use distribution diagnostics",
            "",
            "| Resolution | Source | Author target | Best candidate | RMSE |",
            "|---:|---|---|---|---:|",
        ]
    )
    for item in report["best_landuse_distribution_matches"]:
        lines.append(
            f"| {item['resolution_degrees']:.6f} | "
            f"`{item['landuse_source']}` | `{item['target']}` | "
            f"`{item['candidate']}` | {item['quantile_rmse']:.6f} |"
        )
    lines.extend(["", "## Remaining work", ""])
    lines.extend(f"- {item}" for item in report["remaining_stages"])
    lines.extend(["", "## Methodological limits", ""])
    lines.extend(f"- {item}" for item in report["interpretation"])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def parse_args() -> argparse.Namespace:
    raw_root = default_raw_root()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--boundary",
        type=Path,
        default=raw_root / "boundary" / "communes4326.geojson",
    )
    parser.add_argument(
        "--traffic-workbook",
        type=Path,
        default=raw_root / "traffic" / "donneestrafic-2018.xlsx",
    )
    parser.add_argument(
        "--author-csv",
        type=Path,
        default=(
            Path(__file__).resolve().parents[2]
            / "data"
            / "preprocessed_data"
            / "data_demvk.csv"
        ),
    )
    parser.add_argument(
        "--osm-json",
        type=Path,
        default=raw_root / "osm" / "luxembourg_roads_2021-06-30.json",
    )
    parser.add_argument(
        "--landuse-osm-json",
        type=Path,
        default=raw_root / "osm" / "luxembourg_landuse_ways_2021-06-30.json",
    )
    parser.add_argument(
        "--landuse-geofabrik-zip",
        type=Path,
        default=raw_root / "osm" / "luxembourg-210101-free.shp.zip",
    )
    parser.add_argument(
        "--resolutions",
        type=float,
        nargs="+",
        default=list(DEFAULT_RESOLUTIONS),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("output") / "airpollution_reconstruction" / "stage1",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    report = reconstruct_stage1(
        boundary_path=args.boundary,
        traffic_workbook_path=args.traffic_workbook,
        output_dir=args.output_dir,
        resolutions=tuple(args.resolutions),
        author_csv_path=args.author_csv,
        osm_json_path=args.osm_json,
        landuse_osm_json_path=args.landuse_osm_json,
        landuse_geofabrik_zip_path=args.landuse_geofabrik_zip,
    )
    print(args.output_dir / "stage1_report.json")
    print(args.output_dir / "stage1_report.md")
    best = report["best_author_traffic_distribution_match"]
    if best is not None:
        print(best["output_csv"])


if __name__ == "__main__":
    main()
